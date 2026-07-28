import os
import sys
try:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
except ImportError:
    print("python-pptx is not installed. Please install it using 'pip install python-pptx'")
    sys.exit(1)

def create_ppt():
    prs = Presentation()
    
    # 폰트 및 색상 설정 (Design.md 기반 Deep Indigo 및 Emerald Green)
    title_color = RGBColor(7, 2, 53)     # Deep Indigo (#070235)
    accent_color = RGBColor(16, 185, 129) # Emerald Green (#10b981)
    text_color = RGBColor(71, 70, 79)     # On-Surface-Variant (#47464f)
    
    # 1. Title Slide
    slide_layout = prs.slide_layouts[0] # Title Slide layout
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    
    title.text = "AIoT 스마트 투약 디스펜서\n및 안심 케어 시스템"
    title.text_frame.paragraphs[0].font.color.rgb = title_color
    title.text_frame.paragraphs[0].font.name = 'Malgun Gothic'
    title.text_frame.paragraphs[0].font.bold = True
    title.text_frame.paragraphs[0].font.size = Pt(40)
    
    subtitle.text = "독거노인의 안전한 복약 관리와 응급 상황 대처를 위한 스마트홈 솔루션"
    subtitle.text_frame.paragraphs[0].font.color.rgb = accent_color
    subtitle.text_frame.paragraphs[0].font.name = 'Malgun Gothic'
    subtitle.text_frame.paragraphs[0].font.size = Pt(20)
    
    # helper function to add content slides
    def add_slide(prs, title_text, bullet_points):
        layout = prs.slide_layouts[1] # Title and Content
        slide = prs.slides.add_slide(layout)
        title = slide.shapes.title
        title.text = title_text
        title.text_frame.paragraphs[0].font.color.rgb = title_color
        title.text_frame.paragraphs[0].font.name = 'Malgun Gothic'
        title.text_frame.paragraphs[0].font.bold = True
        
        body = slide.placeholders[1]
        tf = body.text_frame
        tf.clear()
        
        for point in bullet_points:
            p = tf.add_paragraph()
            p.text = point
            p.font.name = 'Malgun Gothic'
            p.font.size = Pt(18)
            p.font.color.rgb = text_color
            p.level = 0
            
            # Sub-bullets
            if point.startswith("  ") or point.startswith("- "):
                p.level = 1
                p.font.size = Pt(16)
                p.text = point.strip(" -")
                
    # 2. Background
    add_slide(prs, "작품 제작 배경 (Background & Need)", [
        "점점 늘어나는 독거노인, 그들의 가장 큰 위협은",
        "  '제때 약을 먹지 못하는 것'과 '홀로 맞는 응급상황'입니다.",
        "",
        "고령화 및 독거노인 증가 현황: 돌봄 인력의 한계",
        "복약 순응도 저하 문제: 건망증 등으로 인한 중복/누락 위험",
        "골든타임 확보의 어려움: 응급 상황 시 자력 신고 한계",
        "",
        "결론: 능동적 복약 유도 및 위급 상황 알림 시스템 필요"
    ])
    
    # 3. Core Features
    add_slide(prs, "핵심 기능 및 서비스 흐름", [
        "하드웨어 제어와 웹 대시보드의 실시간 연동으로 빈틈없는 케어 제공",
        "",
        "Smart Dispenser:",
        "  정밀 모터를 통한 정량 투약, 상태 표시 램프, 부저 알림",
        "Web Dashboard:",
        "  보호자 및 관리자를 위한 실시간 모니터링 및 원격 제어",
        "Motion AI Cam:",
        "  제스처 인식 기반의 일상 제어 및 위급 상황 감지"
    ])
    
    # 4. Scenario 1
    add_slide(prs, "시나리오 1: 정상 복약 관리", [
        "[상황] 지정된 약 복용 시간 도래",
        "",
        "1. 알림 발생:",
        "  LCD 화면 '약 복용 시간입니다' + 부저음 + 초록 램프",
        "2. 사용자 인지 및 행동:",
        "  디스펜서 하단에 손을 가져다 댐 (근접 센서 작동)",
        "3. 투약 및 기록:",
        "  약 1회분 배출 → 웹 대시보드에 즉시 '복용 완료' 업데이트"
    ])

    # 5. Scenario 2
    add_slide(prs, "시나리오 2: 복약 지연 및 경고", [
        "[상황] 지정된 시간이 지나도 약을 복용하지 않은 경우",
        "",
        "1. 경고 발생:",
        "  경고 메세지 + 부저음 패턴 변경 + 빨간 램프 점멸",
        "2. 웹 알림 전송:",
        "  웹 대시보드로 긴급 알림 전송 ('약을 복용하지 않았습니다!')",
        "3. 보호자 개입:",
        "  보호자가 알림을 확인하고 안부 확인 유도"
    ])

    # 6. Scenario 3
    add_slide(prs, "시나리오 3: 일상 제어 및 모션 인식", [
        "[상황] 평상시 조명 제어가 필요한 경우",
        "",
        "1. 제스처 인식:",
        "  카메라를 향해 주먹을 쥐거나 엄지를 치켜드는 모션 취함",
        "2. 동작 수행:",
        "  모션을 인식하여 램프 전원 ON / OFF",
        "",
        "효과: 거동이 불편한 노인도 직관적인 제스처로 조명 제어 가능"
    ])

    # 7. Scenario 4
    add_slide(prs, "시나리오 4: 위급 상황 대응", [
        "[상황] 갑작스러운 쓰러짐 등 위급 상황 발생 시",
        "",
        "1. 위급 상황 감지:",
        "  카메라 모션 인식(쓰러짐/구조요청) 또는 SOS 버튼 작동",
        "2. 즉각 알림:",
        "  웹 대시보드에 최우선 순위 긴급 알람 팝업",
        "3. 외부 연계:",
        "  즉시 응급센터(119) 및 보호자에게 상황 알림 전송"
    ])

    # 8. Web Features
    add_slide(prs, "웹 대시보드 확장 기능", [
        "1. 복약 스케줄러 관리:",
        "  웹에서 직접 시간 설정 및 '약 스스로 넣기' 제어",
        "2. 원격 환경 제어:",
        "  램프 밝기 미세 조절, 배터리 및 약 잔여량 확인",
        "3. AI 케어 챗봇 연동:",
        "  질의응답 및 감정 케어를 돕는 대화형 인터페이스",
        "",
        "UI 디자인: 복잡한 정보 대신 카드 형태의 직관적인 미니멀리스트 디자인"
    ])

    # 9. Future
    add_slide(prs, "추후 활용 방안 및 스마트홈 확장성", [
        "단일 기기를 넘어, 노년층을 위한 종합 스마트홈 허브로 진화",
        "",
        "1. 환경 센서 연동 (Smart Home Hub):",
        "  온/습도 센서 등과 연동하여 실내 환경 최적화",
        "2. 빅데이터 & AI 건강 예측:",
        "  활동량/복약 데이터를 분석하여 건강 이상 징후 감지 및 리포트",
        "3. B2B / B2G 모델 확장:",
        "  지자체 노인 복지 사업 및 요양 병원 연계 통합 관제 시스템"
    ])
    
    ppt_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "SmartHome_Presentation.pptx")
    prs.save(ppt_path)
    print(f"Success! PPT saved to: {ppt_path}")

if __name__ == '__main__':
    create_ppt()
